"""État des lieux d'ENTRÉE — nouveau modèle Word « 2027 » (général et
Notre Maison), 100 % local.

Le modèle 2027 est un formulaire : contrôles de contenu (adresse, date,
dossier, initiales, bailleur, expert, qualité), champs Word (REF, DATE),
tableaux clés / compteurs / attestations avec cases à cocher. La dactylo tape
directement dedans et colle les photos. Ce module applique le rapport de
référence validé par Georges le 29/10/2026 (2026.3384 Wasseiges) :

  * photos des pièces dans une grille sans bordure à 2 colonnes, cases
    ≈ 7,9 × 5,9 cm (8 photos par page), portraits centrés dans la case,
    rotations faites dans Word intégrées aux pixels, doublons signalés ;
  * photo des clés flottante en haut à droite, à côté du titre « Clés,
    compteurs et entretiens », sans masquer l'en-tête ;
  * photos des compteurs sur une ligne sous l'encadré RAPPEL des compteurs ;
  * cases « Repris ci-dessous » / « Voir photos » cochées selon le dossier ;
  * bug du modèle corrigé (signet EDL_bailleur vide → le bailleur
    s'effaçait à la mise à jour des champs) ;
  * initiales, expert et qualité de l'expert ; date du constat sans l'année
    (« 1er octobre ») ; commune en MAJUSCULES ; noms des parties en gras,
    taille par défaut ;
  * date de clôture figée (le champ DATE automatique devient du texte) et
    signature dans la case de l'expert ;
  * modèle Notre Maison : bailleur gardé tel quel, remplissage d'essai du
    modèle (CH, Charline HAZEE, adresse du bureau) remplacé ou signalé.

Le texte dicté n'est JAMAIS modifié.
"""
import io
import os
import re
import tempfile
import zipfile
from copy import deepcopy
from datetime import date

from PIL import Image, ImageOps

import entree_engine as ee
import securite

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
}
W = "{%s}" % NS["w"]
A = "{%s}" % NS["a"]
R = "{%s}" % NS["r"]
WP = "{%s}" % NS["wp"]
W14 = "{%s}" % NS["w14"]
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

EMU_PER_CM = 360000
TWIP_PER_CM = 1440 / 2.54

# Grille des photos (rapport de référence 2026.3384)
GRILLE_CASE_H_CM = 5.9
GRILLE_CASE_W_CM = 7.87
GRILLE_COL_CM = 8.5
GRILLE_ESPACE_CM = 0.25
# Photo des clés (flottante) et des compteurs (une ligne)
CLES_H_CM = 3.0
CLES_W_MAX_CM = 4.0
CLES_TOP_CM = 0.1
COMPTEURS_H_CM = 3.0
COMPTEURS_LARGEUR_MAX_CM = 16.0
# Signature dans la case de l'expert (comme dans le rapport de référence)
SIGNATURE_W_EMU = 1232934
SIGNATURE_POS_H = 1664305
SIGNATURE_POS_V = 112203

COMPRESS_MAX_PX = 1400
COMPRESS_QUALITY = 80

EXPERT_NOM = "Georges BASTIAN"
EXPERT_QUALITE = "Géomètre-expert & Expert immobilier agréé"

ROLES = ("piece", "cles", "compteur", "exclure")
ROLE_LABELS = {"piece": "Pièce", "cles": "Clés", "compteur": "Compteur", "exclure": "Exclure"}

MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre")

NOTRE_MAISON = "NOTRE MAISON"
NM_ESSAI = ("REINE ASTRID", "HAZEE")


class Entree2027Error(Exception):
    pass


# --------------------------------------------------------------------------
# Détection du modèle
# --------------------------------------------------------------------------
def _document_xml(docx_path):
    try:
        with zipfile.ZipFile(docx_path) as z:
            return z.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return ""


def is_model_2027(docx_path, xml=None):
    """True si le Word est le formulaire 2027 (styles « EDL… » et contrôles
    de contenu du modèle)."""
    xml = xml if xml is not None else _document_xml(docx_path)
    return 'w:val="EDLTitredesection"' in xml or ('w:tag w:val="dossier"' in xml and "EDL_" in xml)


def model_kind(docx_path):
    """« 2027-NM » (Notre Maison), « 2027 » (général) ou « ancien »."""
    xml = _document_xml(docx_path)
    if not is_model_2027(docx_path, xml):
        return "ancien"
    plain = ee._norm(re.sub(r"<[^>]+>", "", xml))
    return "2027-NM" if NOTRE_MAISON in plain else "2027"


def needs_finalize(docx_path):
    """True tant que la date de clôture n'est pas figée (le champ DATE
    automatique est remplacé par du texte lors de la finalisation)."""
    xml = _document_xml(docx_path)
    return bool(re.search(r"<w:instrText[^>]*>\s*DATE\b", xml))


# --------------------------------------------------------------------------
# Petits outils XML
# --------------------------------------------------------------------------
def _text(el):
    return "".join(t.text or "" for t in el.iter(W + "t"))


def _children(body):
    return list(body.iterchildren())


def _find_child(els, pred, start=0):
    for i in range(start, len(els)):
        if pred(els[i]):
            return i
    return None


def _is_p(el):
    return el.tag == W + "p"


def _is_tbl(el):
    return el.tag == W + "tbl"


def _has_page_break(el):
    return any(br.get(W + "type") == "page" for br in el.iter(W + "br"))


def _has_drawing(el):
    return el.find(".//" + W + "drawing") is not None


def _sdt_alias(sdt):
    a = sdt.find(W + "sdtPr/" + W + "alias")
    return a.get(W + "val") if a is not None else ""


def _sdt_tag(sdt):
    t = sdt.find(W + "sdtPr/" + W + "tag")
    return t.get(W + "val") if t is not None else ""


def _sdts(root, tag):
    return [s for s in root.iter(W + "sdt") if _sdt_tag(s) == tag]


def _sdt_text(sdt):
    c = sdt.find(W + "sdtContent")
    return _text(c) if c is not None else ""


def _set_sdt_text(sdt, value):
    """Remplace le texte affiché d'un contrôle de contenu (garde la mise en
    forme du premier morceau de texte)."""
    c = sdt.find(W + "sdtContent")
    if c is None:
        return False
    ts = list(c.iter(W + "t"))
    if not ts:
        return False
    ts[0].text = value
    ts[0].set(XML_SPACE, "preserve")
    for t in ts[1:]:
        t.text = ""
    pr = sdt.find(W + "sdtPr")
    if pr is not None:
        ph = pr.find(W + "showingPlcHdr")
        if ph is not None:
            pr.remove(ph)
    return True


def _make(tag, **attrs):
    from docx.oxml import OxmlElement

    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(W + k if ":" not in k else k, v)
    return el


def _rpr(run):
    rpr = run.find(W + "rPr")
    if rpr is None:
        rpr = _make("w:rPr")
        run.insert(0, rpr)
    return rpr


def _highlight(el):
    for r in el.iter(W + "r"):
        if r.find(W + "t") is None:
            continue
        rpr = _rpr(r)
        hl = rpr.find(W + "highlight")
        if hl is None:
            hl = _make("w:highlight")
            rpr.append(hl)
        hl.set(W + "val", "yellow")


def _replace_text(p_el, new_text):
    runs = [r for r in p_el.iter(W + "r") if r.find(W + "t") is not None]
    if not runs:
        return
    t = runs[0].find(W + "t")
    t.text = new_text
    t.set(XML_SPACE, "preserve")
    for r in runs[1:]:
        for t2 in r.iter(W + "t"):
            t2.text = ""


def format_date_cloture(date_cloture):
    """« 29/10/2026 » → « 29 octobre 2026 » (format du champ DATE du
    modèle : d MMMM yyyy)."""
    m = re.match(r"\s*(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", date_cloture or "")
    if not m:
        d = date.today()
        return f"{d.day} {MOIS[d.month - 1]} {d.year}"
    j, mo, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if a < 100:
        a += 2000
    return f"{j} {MOIS[mo - 1]} {a}"


def format_date_constat(value):
    """Date du constat sans l'année (le modèle l'ajoute) : « 01.10. »,
    « 1/10/2026 », « 01-10 » → « 1er octobre ». Autre texte : inchangé."""
    m = re.match(r"^\s*(\d{1,2})\s*[./-]\s*(\d{1,2})\s*(?:[./-]\s*(\d{2,4}))?\s*\.?\s*$", value or "")
    if not m:
        return None
    j, mo = int(m.group(1)), int(m.group(2))
    if not (1 <= j <= 31 and 1 <= mo <= 12):
        return None
    return f"{'1er' if j == 1 else j} {MOIS[mo - 1]}"


def format_localite(value):
    """« 4219 Wasseiges » → « 4219 WASSEIGES » (commune en majuscules)."""
    m = re.match(r"^(\s*\d{4}\s*[,\-]?\s*)(.*)$", value or "")
    if not m or not m.group(2).strip():
        return None
    new = m.group(1) + m.group(2).upper()
    return new if new != value else None


# --------------------------------------------------------------------------
# Champs REF / DATE
# --------------------------------------------------------------------------
def _fields_in_paragraph(p):
    """Champs complexes d'un paragraphe : [{"instr", "runs", "result_ts"}],
    `runs` = toutes les balises <w:r> du champ (début → fin), dans l'ordre."""
    out, stack = [], []
    for el in p.iter():
        if el.tag == W + "fldChar":
            typ = el.get(W + "fldCharType")
            run = el.getparent()
            if typ == "begin":
                stack.append({"instr": "", "runs": [run], "result_ts": [], "in_result": False})
            elif stack:
                f = stack[-1]
                if run not in f["runs"]:
                    f["runs"].append(run)
                if typ == "separate":
                    f["in_result"] = True
                elif typ == "end":
                    out.append(stack.pop())
        elif stack:
            f = stack[-1]
            if el.tag == W + "instrText":
                f["instr"] += el.text or ""
                if el.getparent() not in f["runs"]:
                    f["runs"].append(el.getparent())
            elif el.tag == W + "t" and f["in_result"]:
                f["result_ts"].append(el)
            elif el.tag == W + "r" and el not in f["runs"]:
                f["runs"].append(el)
    return out


def _bookmark_texts(body):
    """Texte contenu dans chaque signet : {nom: texte}."""
    starts = {b.get(W + "id"): b.get(W + "name") for b in body.iter(W + "bookmarkStart")}
    texts, open_ids = {}, set()
    for el in body.iter():
        if el.tag == W + "bookmarkStart":
            open_ids.add(el.get(W + "id"))
            texts.setdefault(starts[el.get(W + "id")], "")
        elif el.tag == W + "bookmarkEnd":
            open_ids.discard(el.get(W + "id"))
        elif el.tag == W + "t" and open_ids:
            for i in open_ids:
                texts[starts[i]] += el.text or ""
    return texts


def _unwrap_field(p, field):
    """Retire la mécanique d'un champ en gardant son résultat affiché comme
    texte normal (début, code et séparateur, fin)."""
    for run in field["runs"]:
        has_result_text = any(t in field["result_ts"] for t in run.iter(W + "t"))
        if has_result_text:
            for fc in list(run.iter(W + "fldChar")):
                fc.getparent().remove(fc)
            continue
        if run.getparent() is not None and (run.find(W + "fldChar") is not None or run.find(W + "instrText") is not None):
            run.getparent().remove(run)


def fix_bailleur_refs(body):
    """Bug du modèle 2027 : le signet EDL_bailleur est vide et le bailleur
    est affiché par un champ REF qui pointe vers lui → toute mise à jour des
    champs (impression, PDF, F9) efface le nom du bailleur. Les REF dont le
    signet est vide deviennent du texte, puis le signet est remis autour du
    contrôle « Bailleur » de la page 2. Retourne le nombre de corrections."""
    texts = _bookmark_texts(body)
    n = 0
    for p in list(body.iter(W + "p")):
        for f in _fields_in_paragraph(p):
            m = re.match(r"\s*REF\s+(\S+)", f["instr"])
            if m and not texts.get(m.group(1), "").strip():
                _unwrap_field(p, f)
                n += 1
    starts = [b for b in body.iter(W + "bookmarkStart") if b.get(W + "name") == "EDL_bailleur"]
    if starts:
        bs = starts[0]
        be = next((e for e in body.iter(W + "bookmarkEnd") if e.get(W + "id") == bs.get(W + "id")), None)
        if be is not None and not _bookmark_texts(body).get("EDL_bailleur", "").strip():
            cands = [s for s in body.iter(W + "sdt") if _sdt_alias(s) == "Bailleur"]
            if cands:
                tgt = cands[-1]
                bs.getparent().remove(bs)
                be.getparent().remove(be)
                tgt.addprevious(bs)
                tgt.addnext(be)
                n += 1
    return n


def sync_ref_results(body, changes):
    """Met à jour le résultat affiché des champs REF dont le texte était une
    ancienne valeur modifiée (`changes` : {ancien: nouveau})."""
    if not changes:
        return
    norm = {k.strip(): v for k, v in changes.items() if k and k.strip()}
    for p in body.iter(W + "p"):
        for f in _fields_in_paragraph(p):
            if not re.match(r"\s*REF\s", f["instr"]) or not f["result_ts"]:
                continue
            cur = "".join(t.text or "" for t in f["result_ts"]).strip()
            if cur in norm:
                f["result_ts"][0].text = norm[cur]
                f["result_ts"][0].set(XML_SPACE, "preserve")
                for t in f["result_ts"][1:]:
                    t.text = ""


def freeze_date_fields(body, texte):
    """Remplace chaque champ DATE automatique par le texte de la date (sinon
    la date change à chaque ouverture du Word). Retourne le nombre."""
    n = 0
    for p in list(body.iter(W + "p")):
        for f in _fields_in_paragraph(p):
            if not re.match(r"\s*DATE\b", f["instr"]):
                continue
            first = f["runs"][0]
            new = deepcopy(first)
            for ch in list(new):
                if ch.tag != W + "rPr":
                    new.remove(ch)
            t = _make("w:t")
            t.text = texte
            t.set(XML_SPACE, "preserve")
            new.append(t)
            first.addprevious(new)
            for run in f["runs"]:
                if run.getparent() is not None:
                    run.getparent().remove(run)
            n += 1
    for fs in list(body.iter(W + "fldSimple")):
        if re.match(r"\s*DATE\b", fs.get(W + "instr") or ""):
            r = _make("w:r")
            t = _make("w:t")
            t.text = texte
            r.append(t)
            fs.addprevious(r)
            fs.getparent().remove(fs)
            n += 1
    return n


# --------------------------------------------------------------------------
# Repères dans le document
# --------------------------------------------------------------------------
def _locate(body):
    els = _children(body)
    norm = lambda el: ee._norm(_text(el))  # noqa: E731
    loc = {"els": els}
    loc["titre_cles"] = _find_child(els, lambda e: _is_p(e) and norm(e).startswith("CLES, COMPTEURS ET ENTRETIENS"))
    loc["rubrique_cles"] = _find_child(els, lambda e: _is_p(e) and norm(e) == "INVENTAIRE DES CLES")
    loc["rubrique_compteurs"] = _find_child(els, lambda e: _is_p(e) and norm(e) == "RELEVE DES COMPTEURS")
    loc["rubrique_attest"] = _find_child(els, lambda e: _is_p(e) and norm(e) == "ATTESTATIONS D'ENTRETIEN")
    loc["rappel_compteurs"] = _find_child(els, lambda e: _is_tbl(e) and "RELEVES NON COMPLETES" in norm(e))
    loc["rappel_attest"] = _find_child(els, lambda e: _is_tbl(e) and "NE CONSERVE NI COPIE" in norm(e))
    loc["titre_photos"] = _find_child(els, lambda e: _is_tbl(e) and "PHOTOGRAPHIES" in norm(e) and "CERTIFI" in norm(e)
                                      and "AUTRES" not in norm(e))
    loc["autres_photos"] = _find_child(els, lambda e: _is_tbl(e) and "AUTRES PHOTOGRAPHIES" in norm(e))
    loc["titre_cloture"] = _find_child(els, lambda e: _is_p(e) and norm(e).startswith("COMPLEMENTS ET CLOTURE"))
    loc["cloture"] = _find_child(els, lambda e: _is_tbl(e) and "EXPERT EN CHARGE" in norm(e) and "FAIT A" in norm(e))
    return loc


def _keys_region(loc):
    """(début, fin) des éléments de la rubrique clés / compteurs /
    attestations où peuvent se trouver les photos des clés et compteurs."""
    i0 = loc["titre_cles"]
    if i0 is None:
        return None
    if i0 > 0 and _is_p(loc["els"][i0 - 1]) and not _text(loc["els"][i0 - 1]).strip():
        i0 -= 1
    i1 = loc["rappel_attest"] if loc["rappel_attest"] is not None else loc["rubrique_attest"]
    if i1 is None:
        return None
    return i0, i1


# --------------------------------------------------------------------------
# Photos déjà collées dans le Word
# --------------------------------------------------------------------------
def _drawing_image(part, drawing):
    """Image d'un dessin Word telle qu'elle s'affiche : recadrage, miroir et
    rotation appliqués aux pixels."""
    blip = drawing.find(".//" + A + "blip")
    if blip is None:
        return None
    rid = blip.get(R + "embed")
    try:
        blob = part.related_parts[rid].blob
    except KeyError:
        return None
    im = Image.open(io.BytesIO(blob))
    im = ImageOps.exif_transpose(im).convert("RGB")
    src = drawing.find(".//" + A + "srcRect")
    if src is not None:
        w, h = im.size
        l, t, r, b = (int(src.get(k, "0")) / 100000 for k in ("l", "t", "r", "b"))
        box = (int(w * l), int(h * t), int(w * (1 - r)), int(h * (1 - b)))
        if box[2] > box[0] and box[3] > box[1]:
            im = im.crop(box)
    xfrm = drawing.find(".//" + A + "xfrm")
    if xfrm is not None:
        if xfrm.get("flipH") == "1":
            im = ImageOps.mirror(im)
        if xfrm.get("flipV") == "1":
            im = ImageOps.flip(im)
        rot = int(xfrm.get("rot", "0") or 0) / 60000
        if rot:
            im = im.rotate(-rot, expand=True, resample=Image.BICUBIC)
    return im


def extract_word_photos(docx_path, out_dir=None, progress_cb=None):
    """Photos déjà collées par la dactylo (rubrique clés/compteurs et section
    PHOTOGRAPHIES), enregistrées en JPEG dans un dossier temporaire (hors
    Dropbox). Retourne [{"path", "role", "note", "rot"}] dans l'ordre du
    Word : photo flottante de la rubrique clés → « clés », autres photos de
    la rubrique → « compteur », section PHOTOGRAPHIES → « pièce »."""
    import docx

    d = docx.Document(docx_path)
    body = d.element.body
    loc = _locate(body)
    out_dir = out_dir or tempfile.mkdtemp(prefix="entree2027_")
    os.makedirs(out_dir, exist_ok=True)
    found = []
    kr = _keys_region(loc)
    if kr:
        for el in loc["els"][kr[0]: kr[1] + 1]:
            for dr in el.iter(W + "drawing"):
                floating = dr.find(WP + "anchor") is not None
                found.append((dr, "cles" if floating else "compteur", "rubrique clés/compteurs du Word"))
    if loc["titre_photos"] is not None:
        end = loc["autres_photos"] if loc["autres_photos"] is not None else len(loc["els"])
        for el in loc["els"][loc["titre_photos"]: end]:
            for dr in el.iter(W + "drawing"):
                found.append((dr, "piece", ""))
    items = []
    for n, (dr, role, note) in enumerate(found, 1):
        if progress_cb:
            progress_cb(f"Lecture des photos du Word {n}/{len(found)}…")
        im = _drawing_image(d.part, dr)
        if im is None:
            continue
        p = os.path.join(out_dir, f"photo_word_{n:03d}.jpg")
        if max(im.size) > 2000:
            im.thumbnail((2000, 2000), Image.LANCZOS)
        im.save(p, "JPEG", quality=92)
        items.append({"path": p, "role": role, "note": note, "rot": 0})
    if kr is None or not any(i["role"] == "cles" for i in items):
        cles = [i for i in items if i["role"] == "compteur"]
        if cles and kr:
            cles[0]["note"] = "à vérifier : clés ou compteur ?"
    return items


def folder_photo_items(folder):
    """Photos du dossier (si la dactylo ne les a pas collées dans le Word)."""
    inv = ee.scan_folder(folder)
    items = []
    for p, checked, note in ee.classify_photos(inv["photo_images"] + inv["root_images"]):
        items.append({"path": p, "role": "piece" if checked else "exclure", "note": note, "rot": 0})
    return items


def _dhash(path, size=8):
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("L").resize((size + 1, size), Image.LANCZOS)
        px = list(im.getdata())
    bits = 0
    for y in range(size):
        for x in range(size):
            bits = (bits << 1) | (px[y * (size + 1) + x] > px[y * (size + 1) + x + 1])
    return bits


def mark_duplicates(items, threshold=12, window=4, exclude=True):
    """Signale les photos quasi identiques à l'une des précédentes (même
    cadrage) et, si `exclude`, les met sur « Exclure ». Les vues proches
    mais différentes (lumière allumée/éteinte) restent : seuil strict."""
    hashes = []
    for it in items:
        try:
            hashes.append(_dhash(it["path"]))
        except Exception:  # noqa: BLE001
            hashes.append(None)
    for i, it in enumerate(items):
        if it["role"] != "piece" or hashes[i] is None:
            continue
        for j in range(max(0, i - window), i):
            if hashes[j] is None or items[j]["role"] == "exclure":
                continue
            if bin(hashes[i] ^ hashes[j]).count("1") <= threshold:
                it["note"] = f"doublon probable de la photo {j + 1}"
                if exclude:
                    it["role"] = "exclure"
                break
    return items


def thumbnail(path, rot=0, size=(150, 112)):
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        if rot:
            im = im.rotate(-rot, expand=True)
        im.thumbnail(size)
        return im.copy()


def _prepare(path, rot=0):
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        if rot:
            im = im.rotate(-rot, expand=True)
        if max(im.size) > COMPRESS_MAX_PX:
            im.thumbnail((COMPRESS_MAX_PX, COMPRESS_MAX_PX), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=COMPRESS_QUALITY, optimize=True)
        return buf.getvalue(), im.width, im.height


# --------------------------------------------------------------------------
# Parties
# --------------------------------------------------------------------------
_LABELS_LOC = ("LOCATAIRE", "LOCATAIRE, D'UNE PART")
_LABELS_BAIL = ("BAILLEUR", "BAILLEUR, D'AUTRE PART")


def _party_name_paragraphs(body, labels):
    """Paragraphes portant le nom d'une partie : le premier paragraphe non
    vide qui suit l'intitulé (page de garde et page 2)."""
    out = []
    for tc in body.iter(W + "tc"):
        ps = list(tc.iter(W + "p"))
        for k, p in enumerate(ps):
            if ee._norm(_text(p)) in labels:
                nxt = next((q for q in ps[k + 1:] if _text(q).strip()), None)
                if nxt is not None:
                    out.append(nxt)
                break
    return out


def _domicile_paragraphs(body, label):
    for tc in body.iter(W + "tc"):
        ps = list(tc.iter(W + "p"))
        if ps and ee._norm(_text(ps[0])) == label:
            return [q for q in ps[2:] if _text(q).strip()]
    return []


def read_parties(docx_path):
    """[{"key", "role", "texte", "editable"}] : noms (modifiables) et
    domiciles (contrôle visuel / surlignage) des parties."""
    import docx

    body = docx.Document(docx_path).element.body
    out = []
    loc = _party_name_paragraphs(body, _LABELS_LOC)
    if loc:
        out.append({"key": "locataire", "role": "Locataire", "texte": _text(loc[0]).strip(), "editable": True})
    bail = _sdts(body, "bailleur") or _party_name_paragraphs(body, _LABELS_BAIL)
    if bail:
        t = _sdt_text(bail[0]) if bail[0].tag == W + "sdt" else _text(bail[0])
        out.append({"key": "bailleur", "role": "Bailleur", "texte": t.strip(), "editable": True})
    for key, role, label in (("locataire_dom", "Dom. locataire", "LOCATAIRE, D'UNE PART"),
                             ("bailleur_dom", "Dom. bailleur", "BAILLEUR, D'AUTRE PART")):
        ps = _domicile_paragraphs(body, label)
        if ps:
            out.append({"key": key, "role": role, "texte": " ".join(_text(p).strip() for p in ps), "editable": False})
    return out


def _apply_parties(body, updates, remarques):
    for upd in updates or []:
        key, texte, hl = upd.get("key"), upd.get("texte"), upd.get("surligner")
        if key == "locataire":
            targets = _party_name_paragraphs(body, _LABELS_LOC)
            for p in targets:
                if texte is not None and texte.strip() and texte.strip() != _text(p).strip():
                    sdt = p.find(".//" + W + "sdt")
                    if sdt is not None and _sdt_text(sdt).strip() == _text(p).strip():
                        _set_sdt_text(sdt, texte.strip())
                    else:
                        _replace_text(p, texte.strip())
                if hl:
                    _highlight(p)
        elif key == "bailleur":
            sdts = _sdts(body, "bailleur")
            if sdts:
                for s in sdts:
                    if texte is not None and texte.strip() and texte.strip() != _sdt_text(s).strip():
                        _set_sdt_text(s, texte.strip())
                    if hl:
                        _highlight(s)
            else:
                for p in _party_name_paragraphs(body, _LABELS_BAIL):
                    if texte is not None and texte.strip() and texte.strip() != _text(p).strip():
                        _replace_text(p, texte.strip())
                    if hl:
                        _highlight(p)
        elif key in ("locataire_dom", "bailleur_dom") and hl:
            label = "LOCATAIRE, D'UNE PART" if key == "locataire_dom" else "BAILLEUR, D'AUTRE PART"
            for p in _domicile_paragraphs(body, label):
                _highlight(p)


def _party_fonts(body):
    """Les 4 noms des parties en gras, taille par défaut (le modèle met le
    locataire en 12 pt)."""
    for p in _party_name_paragraphs(body, _LABELS_LOC) + _party_name_paragraphs(body, _LABELS_BAIL):
        for r in p.iter(W + "r"):
            if r.find(W + "t") is None:
                continue
            rpr = _rpr(r)
            for tag in ("sz", "szCs"):
                el = rpr.find(W + tag)
                if el is not None:
                    rpr.remove(el)
            if rpr.find(W + "b") is None:
                b = _make("w:b")
                # ordre du schéma : rStyle, rFonts, b…
                anchor = rpr.find(W + "rFonts") if rpr.find(W + "rFonts") is not None else rpr.find(W + "rStyle")
                if anchor is not None:
                    anchor.addnext(b)
                else:
                    rpr.insert(0, b)


# --------------------------------------------------------------------------
# Cases à cocher
# --------------------------------------------------------------------------
def _checkboxes(p):
    return [s for s in p.iter(W + "sdt") if s.find(".//" + W14 + "checkbox") is not None]


def _tick(sdt):
    c = sdt.find(".//" + W14 + "checked")
    if c is not None:
        c.set(W14 + "val", "1")
    for t in sdt.iter(W + "t"):
        if t.text == "☐":
            t.text = "☒"


def _table_filled(tbl):
    rows = tbl.findall(W + "tr")
    for tr in rows[1:]:
        cells = tr.findall(W + "tc")
        if any(_text(tc).strip() for tc in cells[1:]):
            return True
    return False


def _section_note_and_table(els, i_rubrique):
    if i_rubrique is None:
        return None, None
    note = _find_child(els, lambda e: _is_p(e) and _checkboxes(e), i_rubrique + 1)
    if note is None or note > i_rubrique + 3:
        return None, None
    tbl = _find_child(els, _is_tbl, note + 1)
    return els[note], (els[tbl] if tbl is not None and tbl <= note + 2 else None)


# --------------------------------------------------------------------------
# Images dans le Word
# --------------------------------------------------------------------------
def _picture_run(document, data, cx, cy):
    from docx.shared import Emu

    tmp = document.add_paragraph()
    run = tmp.add_run()
    run.add_picture(io.BytesIO(data), width=Emu(int(cx)), height=Emu(int(cy)))
    r = run._r
    tmp._p.getparent().remove(tmp._p)
    return r


def _anchor_xml(inline, pos_h, pos_v, wrap, rel_height, dist_l="114300", dist_r="114300"):
    """Transforme une image en ligne en image flottante. `pos_h` / `pos_v` :
    (relativeFrom, "align"|"offset", valeur)."""
    from docx.oxml import parse_xml

    def pos(tag, spec):
        rel, kind, val = spec
        inner = f"<wp:align>{val}</wp:align>" if kind == "align" else f"<wp:posOffset>{int(val)}</wp:posOffset>"
        return f'<wp:{tag} relativeFrom="{rel}">{inner}</wp:{tag}>'

    nswp = f'xmlns:wp="{NS["wp"]}"'
    wrap_xml = {"none": f"<wp:wrapNone {nswp}/>", "square": f'<wp:wrapSquare {nswp} wrapText="bothSides"/>'}[wrap]
    xml = (
        f'<wp:anchor xmlns:wp="{NS["wp"]}" distT="0" distB="0" distL="{dist_l}" distR="{dist_r}" simplePos="0" '
        f'relativeHeight="{rel_height}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>' + pos("positionH", pos_h) + pos("positionV", pos_v) + "</wp:anchor>"
    )
    anchor = parse_xml(xml)
    ext = inline.find(WP + "extent")
    anchor.append(deepcopy(ext))
    eff = inline.find(WP + "effectExtent")
    anchor.append(deepcopy(eff) if eff is not None else parse_xml(
        f'<wp:effectExtent xmlns:wp="{NS["wp"]}" l="0" t="0" r="0" b="0"/>'))
    anchor.append(parse_xml(wrap_xml))
    for tag in ("docPr", "cNvGraphicFramePr"):
        el = inline.find(WP + tag)
        if el is not None:
            anchor.append(deepcopy(el))
    anchor.append(deepcopy(inline.find(A + "graphic")))
    return anchor


def _floating_run(document, data, cx, cy, pos_h, pos_v, wrap, rel_height):
    run = _picture_run(document, data, cx, cy)
    inline = run.find(".//" + WP + "inline")
    anchor = _anchor_xml(inline, pos_h, pos_v, wrap, rel_height)
    inline.getparent().replace(inline, anchor)
    return run


def _drop_unused_rels(document, rids):
    xml = document.element.xml
    for rid in rids:
        if f'"{rid}"' not in xml:
            try:
                document.part.drop_rel(rid)
            except Exception:  # noqa: BLE001
                pass


def _remove_drawings(el, removed_rids):
    for dr in list(el.iter(W + "drawing")):
        blip = dr.find(".//" + A + "blip")
        if blip is not None:
            removed_rids.add(blip.get(R + "embed"))
        run = dr.getparent()
        if run.tag == W + "r" and run.getparent() is not None:
            run.getparent().remove(run)
        elif dr.getparent() is not None:
            dr.getparent().remove(dr)


def _photo_grid(document, photos, progress, col_cm=GRILLE_COL_CM, box_w_cm=GRILLE_CASE_W_CM,
                box_h_cm=GRILLE_CASE_H_CM, gap_cm=GRILLE_ESPACE_CM):
    """Tableau sans bordure à 2 colonnes, hauteur de ligne fixe : 8 photos
    par page, toutes alignées."""
    from docx.oxml import parse_xml
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    nsw = f'xmlns:w="{NS["w"]}"'
    rows = (len(photos) + 1) // 2
    t = document.add_table(rows=rows, cols=2)
    tbl_pr = t._tbl.tblPr
    for x in (
        f'<w:tblBorders {nsw}><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/>'
        '<w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>',
        f'<w:jc {nsw} w:val="center"/>',
        f'<w:tblCellMar {nsw}><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
        '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>',
        f'<w:tblLayout {nsw} w:type="fixed"/>',
    ):
        tbl_pr.append(parse_xml(x))
    style = tbl_pr.find(W + "tblStyle")
    if style is not None:
        tbl_pr.remove(style)
    for gc in t._tbl.tblGrid.findall(W + "gridCol"):
        gc.set(W + "w", str(int(col_cm * TWIP_PER_CM)))
    row_h = int((box_h_cm + gap_cm) * TWIP_PER_CM)
    for row in t.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(parse_xml(f"<w:cantSplit {nsw}/>"))
        tr_pr.append(parse_xml(f'<w:trHeight {nsw} w:val="{row_h}" w:hRule="exact"/>'))
    box_w, box_h = box_w_cm * EMU_PER_CM, box_h_cm * EMU_PER_CM
    for j, it in enumerate(photos):
        progress(f"Photo {j + 1}/{len(photos)} : {os.path.basename(it['path'])}")
        cell = t.cell(j // 2, j % 2)
        cell.width = Cm(col_cm)
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsw} w:val="center"/>'))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Cm(0)
        pf.space_after = Cm(0)
        pf.line_spacing = 1.0
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.right_indent = Cm(0)
        data, w, h = _prepare(it["path"], it.get("rot", 0))
        s = min(box_w / w, box_h / h)
        p._p.append(_picture_run(document, data, w * s, h * s))
    if len(photos) % 2:
        last = t.cell(rows - 1, 1)
        last.width = Cm(col_cm)
    return t._tbl


# --------------------------------------------------------------------------
# Finalisation
# --------------------------------------------------------------------------
def finalize(docx_path, items, date_cloture=None, parties_updates=None, initiales="GB", expert=EXPERT_NOM,
             qualite=EXPERT_QUALITE, rename_to=None, signature_path=None, progress_cb=None):
    """Finalise un Word d'entrée au modèle 2027. `items` : photos dans
    l'ordre, [{"path", "role" (piece/cles/compteur/exclure), "rot"}].
    Écrase le Word d'origine (copie de sauvegarde hors Dropbox d'abord),
    puis le renomme. Retourne {"path", "photos", "remarques",
    "sauvegarde"}."""
    import docx
    from docx.oxml import parse_xml

    def say(msg):
        if progress_cb:
            progress_cb(msg)

    securite.check_dossier_folder(os.path.dirname(docx_path))
    securite.ensure_not_locked(docx_path)
    # Les photos lues dans le Word sont dans un dossier temporaire : on les
    # garde en mémoire avant toute modification.
    say("Sauvegarde de l'original (hors Dropbox)…")
    backup = securite.backup_file(docx_path)
    say("Ouverture du Word…")
    d = docx.Document(docx_path)
    body = d.element.body
    remarques = []
    if not is_model_2027(docx_path):
        raise Entree2027Error("Ce Word n'est pas au nouveau modèle 2027.")
    plain = ee._norm(_text(body))
    notre_maison = NOTRE_MAISON in plain

    # 1. Bug du signet EDL_bailleur (avant toute modification des champs)
    n = fix_bailleur_refs(body)
    if n:
        remarques.append(f"Bug du modèle corrigé ({n} champ(s) REF du bailleur) : le bailleur ne s'effacera plus "
                         "à l'impression. À corriger aussi dans le modèle vierge.")

    # 2. Contrôles de contenu : initiales, expert, qualité, date, commune
    changes = {}
    for s in _sdts(body, "initiales"):
        if _sdt_text(s).strip() != initiales:
            changes[_sdt_text(s).strip()] = initiales
            _set_sdt_text(s, initiales)
    for s in _sdts(body, "expert"):
        if _sdt_text(s).strip() != expert:
            changes[_sdt_text(s).strip()] = expert
            _set_sdt_text(s, expert)
    for s in _sdts(body, "qualite"):
        if _sdt_text(s).strip() != qualite:
            _set_sdt_text(s, qualite)
    dossier = ee.parse_folder_name(os.path.dirname(docx_path))["dossier"]
    for s in _sdts(body, "dossier"):
        cur = _sdt_text(s).strip()
        if dossier and (not cur or cur.startswith("[") or cur.endswith(".0000") or not re.match(r"^\d{4}\.\d{3,4}$", cur)):
            _set_sdt_text(s, dossier)
        elif dossier and cur != dossier:
            remarques.append(f"N° de dossier du Word « {cur} » ≠ dossier « {dossier} » : à vérifier.")
            _highlight(s)
    dates = _sdts(body, "date")
    ref_date = next((_sdt_text(s).strip() for s in dates if _sdt_text(s).strip() and not _sdt_text(s).strip().startswith("[")), "")
    new_date = format_date_constat(ref_date) or ref_date
    if new_date:
        for s in dates:
            cur = _sdt_text(s).strip()
            if cur != new_date:
                changes[cur] = new_date
                _set_sdt_text(s, new_date)
    else:
        remarques.append("Date du constat non remplie dans le Word.")
    for s in _sdts(body, "localite"):
        cur = _sdt_text(s)
        new = format_localite(cur)
        if new:
            changes[cur.strip()] = new.strip()
            _set_sdt_text(s, new)
    rue = next((_sdt_text(s) for s in _sdts(body, "rue")), "")
    if any(k in ee._norm(rue) for k in NM_ESSAI) or "1300 WAVRE" in ee._norm(" ".join(_sdt_text(s) for s in _sdts(body, "localite"))):
        remarques.append("Adresse du bien = adresse d'essai du modèle (bureau de Wavre) : à remplacer par l'adresse réelle "
                         "(surlignée en jaune).")
        for s in _sdts(body, "rue") + _sdts(body, "localite"):
            _highlight(s)
    sync_ref_results(body, changes)

    # 3. Parties : corrections de Georges, puis police des noms
    _apply_parties(body, parties_updates, remarques)
    _party_fonts(body)
    if notre_maison and "HAZEE" in ee._norm(_text(body)):
        remarques.append("Le nom « Charline HAZEE » du remplissage d'essai Notre Maison est encore présent : à vérifier.")

    # 4. Photos
    loc = _locate(body)
    if loc["titre_photos"] is None:
        raise Entree2027Error("Tableau « PHOTOGRAPHIES – Certifiées comme prises par l'expert » introuvable.")
    items = [it for it in items if it.get("role") != "exclure"]
    cles = [it for it in items if it["role"] == "cles"]
    compteurs = [it for it in items if it["role"] == "compteur"]
    pieces = [it for it in items if it["role"] == "piece"]
    if len(cles) > 1:
        remarques.append(f"{len(cles)} photos de clés : la première est placée à côté du titre, les autres en tête "
                         "des photos.")
        pieces = cles[1:] + pieces
        cles = cles[:1]
    say("Préparation des photos…")
    # On lit toutes les images avant de retirer les anciennes du Word.
    prepared = {id(it): _prepare(it["path"], it.get("rot", 0)) for it in cles + compteurs}
    removed_rids = set()

    # 4a. Anciennes photos de la rubrique clés/compteurs
    kr = _keys_region(loc)
    if kr:
        for el in loc["els"][kr[0]: kr[1] + 1]:
            _remove_drawings(el, removed_rids)
    # 4b. Section PHOTOGRAPHIES : on vide tout ce qui est entre le titre et
    # « AUTRES PHOTOGRAPHIES » (photos, lignes vides, consigne du modèle),
    # sauf le saut de page et un éventuel texte (gardé et signalé).
    els = loc["els"]
    i0 = loc["titre_photos"]
    i1 = loc["autres_photos"] if loc["autres_photos"] is not None else i0 + 1
    titre = els[i0]
    for tr in titre.findall(W + "tr"):
        if _has_drawing(tr):
            _remove_drawings(tr, removed_rids)
            titre.remove(tr)
    keep_break = None
    for el in els[i0 + 1: i1]:
        txt = _text(el).strip()
        if _is_p(el) and txt.startswith("[") and "PHOTO" in ee._norm(txt):
            body.remove(el)  # consigne surlignée du modèle
            continue
        if txt:
            remarques.append(f"Texte trouvé dans la section photos, conservé : « {txt[:60]} »")
            continue
        if _is_p(el) and _has_page_break(el) and keep_break is None:
            keep_break = el
            _remove_drawings(el, removed_rids)
            continue
        _remove_drawings(el, removed_rids)
        body.remove(el)
    for el in body.iter(W + "p"):
        txt = _text(el).strip()
        if txt.startswith("[Toutes les photos") or (txt.startswith("[") and "PHOTOS COMPTEURS" in ee._norm(txt)):
            el.getparent().remove(el)
            break
    if keep_break is None and loc["autres_photos"] is not None:
        keep_break = _make("w:p")
        r = _make("w:r")
        r.append(_make("w:br", type="page"))
        keep_break.append(r)
        els[i1].addprevious(keep_break)

    # 4c. Grille des photos des pièces
    if pieces:
        grid = _photo_grid(d, pieces, say)
        titre.addnext(grid)
        sp = _make("w:p")
        sp.append(parse_xml(
            f'<w:pPr xmlns:w="{NS["w"]}"><w:spacing w:before="0" w:after="0" w:line="170" w:lineRule="exact"/></w:pPr>'))
        titre.addnext(sp)
    else:
        remarques.append("Aucune photo des pièces.")

    # 4d. Photo des clés : flottante en haut à droite, à côté du titre
    loc = _locate(body)
    if cles and loc["titre_cles"] is not None:
        data, w, h = prepared[id(cles[0])]
        cy = CLES_H_CM * EMU_PER_CM
        cx = cy * w / h
        if cx > CLES_W_MAX_CM * EMU_PER_CM:
            cx = CLES_W_MAX_CM * EMU_PER_CM
            cy = cx * h / w
        i = loc["titre_cles"]
        host = loc["els"][i - 1] if i > 0 and _is_p(loc["els"][i - 1]) and not _text(loc["els"][i - 1]).strip() \
            and not _has_page_break(loc["els"][i - 1]) else loc["els"][i]
        run = _floating_run(d, data, cx, cy, ("margin", "align", "right"),
                            ("margin", "offset", CLES_TOP_CM * EMU_PER_CM), "none", 251659264)
        host.append(run)
    elif not cles:
        remarques.append("Pas de photo des clés désignée.")

    # 4e. Photos des compteurs : une ligne sous l'encadré RAPPEL des compteurs
    if compteurs and loc["rappel_compteurs"] is not None:
        sizes = []
        for it in compteurs:
            _data, w, h = prepared[id(it)]
            sizes.append((COMPTEURS_H_CM * EMU_PER_CM * w / h, COMPTEURS_H_CM * EMU_PER_CM))
        total = sum(cx for cx, _cy in sizes) + 0.4 * EMU_PER_CM * (len(sizes) - 1)
        k = min(1.0, COMPTEURS_LARGEUR_MAX_CM * EMU_PER_CM / total)
        p = _make("w:p")
        p.append(parse_xml(f'<w:pPr xmlns:w="{NS["w"]}"><w:keepLines/><w:spacing w:before="57" w:after="0"/>'
                           '<w:jc w:val="center"/></w:pPr>'))
        for n, (it, (cx, cy)) in enumerate(zip(compteurs, sizes)):
            if n:
                r = _make("w:r")
                t = _make("w:t")
                t.text = "  "
                t.set(XML_SPACE, "preserve")
                r.append(t)
                p.append(r)
            p.append(_picture_run(d, prepared[id(it)][0], cx * k, cy * k))
        loc["els"][loc["rappel_compteurs"]].addnext(p)
    # Page blanche évitée : quand la page clés/compteurs est pleine, le saut
    # de page manuel suivi d'une ligne vide déborde sur une page vide. Le
    # saut manuel est remplacé par « saut de page avant » sur le titre
    # « Informations complémentaires ».
    loc = _locate(body)
    if loc["rappel_attest"] is not None:
        els = loc["els"]
        i_info = _find_child(els, lambda e: _is_p(e) and ee._norm(_text(e)).startswith("INFORMATIONS COMPLEMENTAIRES"),
                             loc["rappel_attest"] + 1)
        between = els[loc["rappel_attest"] + 1: i_info] if i_info is not None else []
        if i_info is not None and all(_is_p(e) and not _text(e).strip() and not _has_drawing(e)
                                      and e.find(".//" + W + "sectPr") is None for e in between):
            for e in between:
                body.remove(e)
            ee._set_ppr_flag(els[i_info], "pageBreakBefore", True)

    # 4f. Cases à cocher
    loc = _locate(body)
    els = loc["els"]
    for rub, photos_ok, label in ((loc["rubrique_cles"], bool(cles), "clés"),
                                  (loc["rubrique_compteurs"], bool(compteurs), "compteurs"),
                                  (loc["rubrique_attest"], False, "attestations")):
        note, tbl = _section_note_and_table(els, rub)
        if note is None:
            continue
        cbs = _checkboxes(note)
        if tbl is not None and _table_filled(tbl) and cbs:
            _tick(cbs[0])
        if photos_ok and len(cbs) > 1:
            _tick(cbs[1])
        if tbl is not None and not _table_filled(tbl) and label != "attestations":
            remarques.append(f"Tableau des {label} vide : vérifier la dictée (relevés entre parties ?).")

    # 5. Saut de page avant « Compléments et clôture » retiré (les autres
    # photographies et la clôture partagent la page, comme le 2026.3384).
    if loc["autres_photos"] is not None and loc["titre_cloture"] is not None:
        for el in els[loc["autres_photos"] + 1: loc["titre_cloture"]]:
            if _is_p(el) and _has_page_break(el) and not _text(el).strip() and not _has_drawing(el):
                body.remove(el)

    # 6. Date de clôture figée
    texte_date = format_date_cloture(date_cloture or date.today().strftime("%d/%m/%Y"))
    if not freeze_date_fields(body, texte_date):
        remarques.append("Champ DATE de la clôture introuvable : date à vérifier.")

    # 7. Signature dans la case de l'expert
    if signature_path is None:
        import signature as _sig

        signature_path = _sig.signature_path() or ""
    loc = _locate(body)
    if loc["cloture"] is not None and signature_path and os.path.exists(signature_path):
        clo = loc["els"][loc["cloture"]]
        target = next((p for p in clo.iter(W + "p") if ee._norm(_text(p)) == "L'EXPERT EN CHARGE"), None)
        tc = target.getparent() if target is not None else None
        if tc is not None and not _has_drawing(tc):
            with Image.open(signature_path) as sim:
                sw, sh = sim.size
            buf = io.BytesIO()
            with open(signature_path, "rb") as f:
                buf.write(f.read())
            cx = SIGNATURE_W_EMU
            cy = cx * sh / sw
            run = _floating_run(d, buf.getvalue(), cx, cy, ("margin", "offset", SIGNATURE_POS_H),
                                ("paragraph", "offset", SIGNATURE_POS_V), "square", 251660288)
            first_run = target.find(W + "r")
            if first_run is not None:
                first_run.addprevious(run)
            else:
                target.append(run)
            # Espace de signature ramené à deux lignes pointillées (comme le
            # 2026.3384) : sinon la clôture déborde d'une ligne sur une page.
            ps = list(tc.findall(W + "p"))
            k = ps.index(target)
            vides = []
            for p in ps[k + 1:]:
                if _text(p).strip() or _has_drawing(p):
                    break
                vides.append(p)
            if len(vides) > 2:
                modele = next((p for p in vides if p.find(".//" + W + "top") is not None), None)
                for p in vides[2:]:
                    tc.remove(p)
                if modele is not None:
                    for p in vides[:2]:
                        if p is not modele:
                            p.getparent().replace(p, deepcopy(modele))
    elif not signature_path:
        remarques.append("Aucune image de signature choisie (menu Aide → Ma signature…) : signature non ajoutée.")
    # Lignes vides après le tableau de clôture : une seule gardée (Word en
    # exige une après un tableau final), sinon page blanche en fin de rapport.
    loc = _locate(body)
    if loc["cloture"] is not None:
        apres = [e for e in loc["els"][loc["cloture"] + 1:] if _is_p(e)]
        if apres and all(not _text(e).strip() and not _has_drawing(e) and not _has_page_break(e)
                         and e.find(".//" + W + "sectPr") is None for e in apres):
            for e in apres[1:]:
                body.remove(e)

    # 8. Repères du modèle restés tels quels (à trancher par Georges)
    restes = sorted({m.group(0) for m in re.finditer(r"\[[^\]\[]{3,80}\]", _text(body))})
    if restes:
        remarques.append("Options / repères du modèle encore présents : " + " ; ".join(restes[:8])
                         + (" …" if len(restes) > 8 else ""))

    _drop_unused_rels(d, removed_rids)

    say("Enregistrement du Word…")
    temp = securite.temp_path_for(docx_path)
    try:
        d.save(temp)
        docx.Document(temp)
    except Exception as exc:
        if os.path.exists(temp):
            os.remove(temp)  # fichier temporaire du programme uniquement
        raise Entree2027Error(f"Enregistrement impossible, l'original n'a pas été modifié : {exc}") from exc
    securite.commit_temp(temp, docx_path)

    final_path = docx_path
    if rename_to:
        target = os.path.join(os.path.dirname(docx_path), rename_to)
        if os.path.normcase(os.path.abspath(target)) != os.path.normcase(os.path.abspath(docx_path)):
            if os.path.exists(target):
                remarques.append(f"« {rename_to} » existait déjà : le Word n'a pas été renommé.")
            else:
                os.replace(docx_path, target)
                final_path = target
    nb = len(pieces) + len(cles) + len(compteurs)
    return {"path": final_path, "photos": nb, "remarques": remarques, "sauvegarde": backup,
            "notre_maison": notre_maison}


# --------------------------------------------------------------------------
# Traitement automatique (lots)
# --------------------------------------------------------------------------
def auto_items(folder, docx_path, progress_cb=None):
    """Photos à utiliser sans intervention : celles du Word (rôles déduits
    de leur place), sinon celles du dossier ; doublons exclus."""
    items = extract_word_photos(docx_path, progress_cb=progress_cb)
    if not items:
        items = folder_photo_items(folder)
    return mark_duplicates(items)


def auto_parties(folder, docx_path, progress_cb=None):
    """Noms des parties du Word comparés aux cartes d'identité du bail
    (sans rien réécrire) ; différence → surlignage."""
    inv = ee.scan_folder(folder)
    names, constats = [], []
    for doc in inv["documents"]:
        if "BAIL" not in ee._norm(os.path.basename(doc)):
            continue
        try:
            r = ee.read_document(doc, progress_cb=progress_cb)
        except Exception as exc:  # noqa: BLE001
            constats.append(f"{os.path.basename(doc)} illisible ({exc})")
            continue
        names += [n for n in r["mrz_names"] if n not in names]
    updates = []
    parties = [p for p in read_parties(docx_path) if p["editable"]]
    if not names:
        constats.append("Pas de nom lu sur des cartes d'identité : parties à contrôler à la main.")
        return updates, constats
    for part in parties:
        if part["key"] == "bailleur" and NOTRE_MAISON in ee._norm(part["texte"]):
            continue
        name, ratio = ee.compare_names(part["texte"], names)
        if ratio >= 0.99:
            constats.append(f"{part['role']} : identique à la carte d'identité ({name}).")
        elif ratio >= 0.6:
            updates.append({"key": part["key"], "texte": None, "surligner": True})
            constats.append(f"{part['role']} : « {part['texte'][:50]} » ≠ carte « {name} » → surligné en jaune.")
        else:
            constats.append(f"{part['role']} : aucun nom correspondant sur les cartes → à contrôler.")
    return updates, constats
