"""État des lieux d'ENTRÉE — finalisation du Word dactylographié, 100 % local.

Le bureau de dactylographie renvoie un Word à finaliser. Ce module :
  * range les photos en vrac dans un sous-dossier « Photos » ;
  * propose la sélection des photos (la capture d'écran du n° de dossier et
    les photos d'index compteurs sont à décocher ; la photo des clés reste) ;
  * injecte les photos sous « PHOTOGRAPHIES - (CERTIFIES COMME PRISES PAR
    L'EXPERT) », titre en haut d'une nouvelle page, 2 photos par page, dans
    l'ordre, avec compression légère (qualité visuelle inchangée) ;
  * place la fin du rapport (« AUTRES PHOTOGRAPHIES… » → signature) sur une
    page ;
  * met la date de clôture du jour et la signature (comme le 2026.3115) ;
  * surligne en jaune les éléments signalés (coordonnées des parties qui
    diffèrent du bail / de l'OM) ;
  * écrase le Word d'origine et le renomme « 2026.xxxx ELLE Commune Rue n° GB.docx ».

Le contenu dicté n'est JAMAIS modifié (les repères du type « 12*42 » restent).

Lecture du bail et de l'OM (sans IA, sans internet) :
  * PDF avec texte (bail généré par une agence) → texte lu directement ;
  * PDF scanné → reconnaissance de caractères Tesseract (gratuit, optionnel),
    utilisée notamment pour lire la zone codée (MRZ) des cartes d'identité
    jointes au bail : noms exacts des parties, à comparer au Word.
"""
import difflib
import io
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from copy import deepcopy
from datetime import date

from PIL import Image, ImageOps

import securite

APP_DIR = os.path.dirname(os.path.abspath(__file__))

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".heic", ".bmp", ".tif", ".tiff", ".webp"}
DOC_EXT = {".pdf", ".docx", ".jpg", ".jpeg", ".png"}

# Taille des photos : 2 par page A4 sous le titre (15,3 × 11,5 cm en paysage)
EMU_PER_CM = 360000
PHOTO_MAX_W = int(15.3 * EMU_PER_CM)
PHOTO_MAX_H = int(11.5 * EMU_PER_CM)
SIGNATURE_W = 1619250  # 4,5 cm, comme dans le 2026.3115
SIGNATURE_H = 809625
COMPRESS_MAX_PX = 2000
COMPRESS_QUALITY = 85

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
W = "{%s}" % NS["w"]


class EntreeError(Exception):
    pass


def _norm(s):
    s = (s or "").replace("\u2019", "'").replace("\u2018", "'").replace("\u00a0", " ")
    s = unicodedata.normalize("NFKD", s)
    s = s.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip().upper()


# --------------------------------------------------------------------------
# Dossier
# --------------------------------------------------------------------------
_DATE_SUFFIX_RE = re.compile(r"\s*\(\s*\d{1,2}[.\-/]\d{1,2}\s*\)\s*")


def parse_folder_name(folder):
    name = os.path.basename(os.path.normpath(folder or ""))
    m = re.match(r"^(\d{4}\.\d{3,4})\s+(\S+)\s+(.*)$", name)
    if not m:
        return {"dossier": "", "type": "", "reste": name, "initiales": ""}
    dossier, typ, reste = m.groups()
    initiales = ""
    m2 = re.search(r"(?:^|[\s)])\(?([A-Za-z]{2,4})\)?\s*$", reste)
    if m2:
        initiales = m2.group(1).upper()
        reste = reste[: m2.start(1)].rstrip(" (")
    reste = _DATE_SUFFIX_RE.sub(" ", reste).strip()
    reste = re.sub(r"\s+", " ", reste)
    return {"dossier": dossier, "type": typ, "reste": reste, "initiales": initiales}


_INITIALES_RE = re.compile(r"(?:^|[\s)])\(?([A-Za-z]{2,4})\)?\s*$")


def folder_initials(folder):
    """Initiales de l'expert en fin de nom de dossier, en majuscules :
    « … (01.09) GB » → GB, « … (05.10)GB » → GB, « … (15.09) bmi » → BMI,
    « … rue du Fort 10 (MN) » → MN. Chaîne vide si absentes."""
    name = os.path.basename(os.path.normpath(folder or "")).strip()
    m = _INITIALES_RE.search(name)
    return m.group(1).upper() if m else ""


def final_docx_name(folder, initiales="GB"):
    """« 2026.3115 ELLE Opprebais Rue d'En Haut 13 (01.09) GB » →
    « 2026.3115 ELLE Opprebais Rue d'En Haut 13 GB.docx »."""
    info = parse_folder_name(folder)
    if not info["dossier"]:
        return None
    ini = info["initiales"] or initiales or "GB"
    name = f"{info['dossier']} {info['type'] or 'ELLE'} {info['reste']} {ini}"
    return re.sub(r'[\\/:*?"<>|]+', "-", name).strip() + ".docx"


def scan_folder(folder):
    """Inventaire du dossier : Word candidats, photos (racine + Photos/),
    documents bail / OM / convention."""
    if not os.path.isdir(folder):
        raise EntreeError(f"Dossier introuvable : {folder}")
    docx_files, root_images, docs = [], [], []
    for f in sorted(os.listdir(folder)):
        p = os.path.join(folder, f)
        if not os.path.isfile(p) or f.startswith("~$"):
            continue
        ext = os.path.splitext(f)[1].lower()
        low = _norm(f)
        if ext == ".docx":
            docx_files.append(p)
        if ext in IMAGE_EXT and not re.search(r"\b(BAIL|OM|ORDRE|CONVENTION|NUMERISATION|SCAN)\b", low):
            root_images.append(p)
        if ext in DOC_EXT and re.search(r"BAIL|\bOM\b|ORDRE DE MISSION|CONVENTION|NUMERISATION|CONTRAT", low):
            docs.append(p)
    photos_dir = os.path.join(folder, "Photos")
    photo_images = []
    if os.path.isdir(photos_dir):
        photo_images = [
            os.path.join(photos_dir, f) for f in sorted(os.listdir(photos_dir))
            if os.path.splitext(f)[1].lower() in IMAGE_EXT and not f.startswith(".")
        ]
    return {"docx": docx_files, "root_images": root_images, "photo_images": photo_images,
            "documents": docs, "photos_dir": photos_dir}


def ranger_photos(folder):
    """Déplace les photos en vrac à la racine dans « Photos ». Ne supprime
    rien : en cas de doublon de nom, la photo est renommée (suffixe _2)."""
    securite.check_dossier_folder(folder)
    inv = scan_folder(folder)
    dest = inv["photos_dir"]
    os.makedirs(dest, exist_ok=True)
    moved = []
    for p in inv["root_images"]:
        target = os.path.join(dest, os.path.basename(p))
        base, ext = os.path.splitext(target)
        k = 2
        while os.path.exists(target):
            target = f"{base}_{k}{ext}"
            k += 1
        shutil.move(p, target)
        moved.append(target)
    return moved


def _num_photo(path):
    m = re.search(r"(\d{3,})", os.path.basename(path))
    return int(m.group(1)) if m else None


def _prefix_photo(path):
    return re.sub(r"\d+.*$", "", os.path.splitext(os.path.basename(path))[0]).upper()


def sort_key(path):
    n = _num_photo(path)
    return (_prefix_photo(path), n if n is not None else 10**9, os.path.basename(path).lower())


def classify_photos(paths):
    """Retourne [(chemin, coché_par_défaut, remarque)] dans l'ordre.
    - captures d'écran (PNG) : décochées (capture du n° de dossier) ;
    - photos d'une autre série (préfixe/numérotation différente de la série
      principale) : cochées mais signalées « autre source ? »."""
    paths = sorted(paths, key=sort_key)
    prefixes = {}
    for p in paths:
        prefixes[_prefix_photo(p)] = prefixes.get(_prefix_photo(p), 0) + 1
    main_prefix = max(prefixes, key=prefixes.get) if prefixes else ""
    out = []
    for p in paths:
        ext = os.path.splitext(p)[1].lower()
        note, checked = "", True
        if ext == ".png" or "CAPTURE" in _norm(os.path.basename(p)) or "SCREENSHOT" in _norm(os.path.basename(p)):
            checked, note = False, "capture d'écran (n° de dossier ?)"
        elif _prefix_photo(p) != main_prefix:
            note = "autre série — photo envoyée par une partie ?"
        out.append((p, checked, note))
    return out


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------
def prepare_image(path):
    """Ouvre la photo (orientation EXIF corrigée), la réduit à 2000 px max
    et la compresse en JPEG qualité 85 (différence invisible à l'impression).
    Retourne (bytes_jpeg, largeur_px, hauteur_px)."""
    try:
        im = Image.open(path)
    except Exception as exc:  # noqa: BLE001
        raise EntreeError(f"Photo illisible : {os.path.basename(path)} ({exc})") from exc
    with im:
        im = ImageOps.exif_transpose(im)
        im = im.convert("RGB")
        if max(im.size) > COMPRESS_MAX_PX:
            im.thumbnail((COMPRESS_MAX_PX, COMPRESS_MAX_PX), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=COMPRESS_QUALITY, optimize=True)
        return buf.getvalue(), im.width, im.height


def fit_box(w, h, max_w=PHOTO_MAX_W, max_h=PHOTO_MAX_H):
    if not w or not h:
        return max_w, max_h
    scale = min(max_w / w, max_h / h)
    return int(w * scale), int(h * scale)


def thumbnail(path, size=(150, 112)):
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        im = im.convert("RGB")
        im.thumbnail(size)
        return im.copy()


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def _ptext(p_el):
    return "".join(t.text or "" for t in p_el.iter(W + "t"))


def _has_image(p_el):
    return bool(p_el.findall(".//{%s}blip" % NS["a"]))


def _set_ppr_flag(p_el, tag, on=True):
    from docx.oxml import OxmlElement

    ppr = p_el.find(W + "pPr")
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    existing = ppr.find(W + tag)
    if on and existing is None:
        el = OxmlElement("w:" + tag)
        # ordre du schéma : keepNext, keepLines, pageBreakBefore viennent
        # juste après pStyle s'il existe
        pstyle = ppr.find(W + "pStyle")
        if pstyle is not None:
            pstyle.addnext(el)
        else:
            ppr.insert(0, el)
    elif not on and existing is not None:
        ppr.remove(existing)


def _remove_page_breaks(p_el):
    removed = 0
    for br in list(p_el.iter(W + "br")):
        if br.get(W + "type") == "page":
            br.getparent().remove(br)
            removed += 1
    return removed


def _find_para(paras, *needles, start=0):
    for i in range(start, len(paras)):
        t = _norm(_ptext(paras[i]))
        if all(_norm(n) in t for n in needles):
            return i
    return None


def read_parties(docx_path):
    """Retourne les paragraphes des parties : [{"role", "index", "texte"}]
    (paragraphe juste avant « Locataire, d'une part » / « Bailleur … d'autre
    part »)."""
    import docx

    d = docx.Document(docx_path)
    paras = [p for p in d.element.body.iterchildren() if p.tag == W + "p"]
    out = []
    for role, needle in (("Locataire", "D'UNE PART"), ("Bailleur", "D'AUTRE PART")):
        i = _find_para(paras, needle)
        if i is None:
            continue
        j = i - 1
        while j >= 0 and not _ptext(paras[j]).strip():
            j -= 1
        if j >= 0:
            out.append({"role": role, "index": j, "texte": _ptext(paras[j]).strip()})
    return out


def _replace_paragraph_text(p_el, new_text):
    """Remplace le texte d'un paragraphe en gardant la mise en forme du
    premier morceau de texte."""
    runs = [r for r in p_el.iter(W + "r") if r.find(W + "t") is not None]
    if not runs:
        return
    first = runs[0]
    first.find(W + "t").text = new_text
    first.find(W + "t").set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for r in runs[1:]:
        r.getparent().remove(r)


def _highlight_paragraph(p_el):
    from docx.oxml import OxmlElement

    for r in p_el.iter(W + "r"):
        if r.find(W + "t") is None:
            continue
        rpr = r.find(W + "rPr")
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        hl = rpr.find(W + "highlight")
        if hl is None:
            hl = OxmlElement("w:highlight")
            rpr.append(hl)
        hl.set(W + "val", "yellow")


def _make_picture_run(document, image_bytes, cx, cy, ref_rpr=None):
    """Crée un <w:r> contenant une image en ligne (taille en EMU)."""
    from docx.shared import Emu

    tmp_par = document.add_paragraph()
    run = tmp_par.add_run()
    run.add_picture(io.BytesIO(image_bytes), width=Emu(cx), height=Emu(cy))
    r_el = run._r
    if ref_rpr is not None:
        r_el.insert(0, deepcopy(ref_rpr))
    tmp_par._p.getparent().remove(tmp_par._p)
    return r_el


def finalize(docx_path, photos, date_cloture=None, parties_updates=None, initiales="GB",
             rename_to=None, signature_path=None, tighten_layout=True, progress_cb=None):
    """Finalise le Word d'entrée. `photos` : chemins dans l'ordre voulu.
    `parties_updates` : [{"index", "texte", "surligner"}] issus de l'écran de
    vérification. Écrase le Word d'origine, puis le renomme. Retourne un
    dict {"path", "photos", "remarques"}."""
    import docx
    from docx.oxml import OxmlElement

    def say(msg):
        if progress_cb:
            progress_cb(msg)

    securite.check_dossier_folder(os.path.dirname(docx_path))
    securite.ensure_not_locked(docx_path)
    say("Sauvegarde de l'original (hors Dropbox)…")
    backup = securite.backup_file(docx_path)

    say("Ouverture du Word…")
    d = docx.Document(docx_path)
    body = d.element.body
    remarques = []

    def paragraphs():
        return [p for p in body.iterchildren() if p.tag == W + "p"]

    # 1. Coordonnées des parties (vérifiées par Georges)
    if parties_updates:
        paras = paragraphs()
        for upd in parties_updates:
            i = upd.get("index")
            if i is None or i >= len(paras):
                continue
            if upd.get("texte") is not None and upd["texte"].strip() != _ptext(paras[i]).strip():
                _replace_paragraph_text(paras[i], upd["texte"])
            if upd.get("surligner"):
                _highlight_paragraph(paras[i])

    # 2. Photos
    paras = paragraphs()
    i_title = _find_para(paras, "PHOTOGRAPHIES", "CERTIFI")
    if i_title is None:
        raise EntreeError("Titre « PHOTOGRAPHIES - (CERTIFIES COMME PRISES PAR L'EXPERT) » introuvable dans le Word.")
    title = paras[i_title]
    i_autres = _find_para(paras, "AUTRES PHOTOGRAPHIES", start=i_title + 1)

    # Titre en haut d'une nouvelle page, sur une ligne
    _set_ppr_flag(title, "pageBreakBefore", True)
    _set_ppr_flag(title, "keepNext", True)
    _remove_page_breaks(title)
    # Saut de page manuel juste avant le titre → inutile (page blanche)
    k = i_title - 1
    while k >= 0 and not _ptext(paras[k]).strip() and not _has_image(paras[k]):
        _remove_page_breaks(paras[k])
        k -= 1

    # Retirer ce qui se trouve entre le titre et « AUTRES PHOTOGRAPHIES »
    # (paragraphes vides ou anciennes photos) — jamais du texte dicté.
    end = i_autres if i_autres is not None else i_title + 1
    ref_ppr, ref_rpr = None, None
    for p in paras[i_title + 1 : end]:
        if _ptext(p).strip():
            remarques.append("Texte trouvé entre le titre des photos et « AUTRES PHOTOGRAPHIES » : conservé.")
            continue
        if ref_ppr is None and p.find(W + "pPr") is not None:
            ref_ppr = deepcopy(p.find(W + "pPr"))
        p.getparent().remove(p)

    photo_par = OxmlElement("w:p")
    ppr = ref_ppr if ref_ppr is not None else OxmlElement("w:pPr")
    if ppr.find(W + "jc") is None:
        jc = OxmlElement("w:jc")
        jc.set(W + "val", "center")
        ppr.append(jc)
    for tag in ("pageBreakBefore", "keepNext"):
        el = ppr.find(W + tag)
        if el is not None:
            ppr.remove(el)
    photo_par.append(ppr)
    title.addnext(photo_par)

    for n, path in enumerate(photos, 1):
        say(f"Photo {n}/{len(photos)} : {os.path.basename(path)}")
        data, wpx, hpx = prepare_image(path)
        cx, cy = fit_box(wpx, hpx)
        photo_par.append(_make_picture_run(d, data, cx, cy))

    # Fin du rapport sur une page : « AUTRES PHOTOGRAPHIES » commence une page
    paras = paragraphs()
    i_autres = _find_para(paras, "AUTRES PHOTOGRAPHIES")
    if i_autres is not None:
        _set_ppr_flag(paras[i_autres], "pageBreakBefore", True)
        _set_ppr_flag(paras[i_autres], "keepNext", True)
        for p in paras[i_autres:]:
            _remove_page_breaks(p)
        # Pour que la fin tienne sur une page : retirer les lignes vides
        # superflues du bloc final (on garde l'espace avant la date de
        # clôture et avant « Po l'expert en charge »).
        keep_before = ("POUR LES PARTIES SUIVANT", "EXPERT EN CHARGE")
        block = paras[i_autres + 1 :]
        for idx, p in enumerate(block):
            if _ptext(p).strip() or _has_image(p) or p.find(".//" + W + "sectPr") is not None:
                continue
            nxt = next((q for q in block[idx + 1 :] if _ptext(q).strip() or _has_image(q)), None)
            if nxt is not None and any(k in _norm(_ptext(nxt)) for k in keep_before):
                continue
            p.getparent().remove(p)

    # 3. Date de clôture
    date_cloture = date_cloture or date.today().strftime("%d/%m/%Y")
    paras = paragraphs()
    i_date = _find_para(paras, "POUR LES PARTIES SUIVANT")
    if i_date is not None:
        p = paras[i_date]
        runs = [r for r in p.iter(W + "r")]
        txt_runs = [r for r in runs if r.find(W + "t") is not None]
        full = "".join(r.find(W + "t").text or "" for r in txt_runs)
        m = re.search(r"\bLe\b", full[full.find(":") + 1 :] if ":" in full else full)
        if m:
            # Reconstruit « … : <tab>Le jj/mm/aaaa » sans toucher au début
            cut = (full.find(":") + 1 if ":" in full else 0) + m.start()
            pos = 0
            done = False
            for r in txt_runs:
                t = r.find(W + "t")
                s = t.text or ""
                if done:
                    t.text = ""
                elif pos + len(s) > cut:
                    t.text = s[: cut - pos] + "Le " + date_cloture
                    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    done = True
                pos += len(s)
        else:
            last = txt_runs[-1] if txt_runs else None
            if last is not None:
                tab = OxmlElement("w:tab")
                last.append(tab)
                t = OxmlElement("w:t")
                t.text = "Le " + date_cloture
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                last.append(t)
    else:
        remarques.append("Ligne « Pour les parties suivant demande bail / ordre de mission » introuvable : date de clôture non mise.")

    # 4. Signature sous « Po l'expert en charge » (sans page en trop)
    if signature_path is None:
        import signature as _sig

        signature_path = _sig.signature_path() or ""
    paras = paragraphs()
    i_po = _find_para(paras, "EXPERT EN CHARGE")
    if i_po is not None and os.path.exists(signature_path):
        nxt = paras[i_po + 1] if i_po + 1 < len(paras) else None
        if nxt is not None and _has_image(nxt):
            pass  # signature déjà présente
        else:
            with open(signature_path, "rb") as f:
                sig = f.read()
            sig_par = OxmlElement("w:p")
            sppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(W + "left", "3824")
            ind.set(W + "firstLine", "424")
            sppr.append(ind)
            sig_par.append(sppr)
            sig_par.append(_make_picture_run(d, sig, SIGNATURE_W, SIGNATURE_H))
            if nxt is not None and not _ptext(nxt).strip() and not _has_image(nxt):
                nxt.addprevious(sig_par)
                nxt.getparent().remove(nxt)  # remplace la ligne vide prévue pour signer
            else:
                paras[i_po].addnext(sig_par)
    elif i_po is None:
        remarques.append("« Po l'expert en charge » introuvable : signature non ajoutée.")
    elif not signature_path:
        remarques.append("Aucune image de signature choisie (menu Aide → Ma signature…) : signature non ajoutée.")

    # 5. Mise en page : lignes vides en série (≥ 3) ramenées à une seule,
    #    hors section photos. Le texte dicté n'est pas touché.
    if tighten_layout:
        paras = paragraphs()
        run = []
        for p in paras + [None]:
            empty = p is not None and not _ptext(p).strip() and not _has_image(p) \
                and p.find(".//" + W + "sectPr") is None and not p.findall(".//" + W + "br")
            if empty:
                run.append(p)
                continue
            if len(run) >= 3:
                for q in run[1:]:
                    q.getparent().remove(q)
            run = []

    say("Enregistrement du Word…")
    temp = securite.temp_path_for(docx_path)
    try:
        d.save(temp)
        import docx as _docx_check

        _docx_check.Document(temp)  # le fichier écrit se relit correctement
    except Exception as exc:
        if os.path.exists(temp):
            os.remove(temp)  # fichier temporaire du programme uniquement
        raise EntreeError(f"Enregistrement impossible, l'original n'a pas été modifié : {exc}") from exc
    securite.commit_temp(temp, docx_path)

    final_path = docx_path
    if rename_to:
        target = os.path.join(os.path.dirname(docx_path), rename_to)
        if os.path.normcase(os.path.abspath(target)) != os.path.normcase(os.path.abspath(docx_path)):
            if os.path.exists(target):
                remarques.append(f"« {rename_to} » existait déjà : le Word n'a pas été renommé.")
            else:
                os.replace(docx_path, target)
                final_path = target
    return {"path": final_path, "photos": len(photos), "remarques": remarques, "sauvegarde": backup}


# --------------------------------------------------------------------------
# Lecture du bail / de l'OM
# --------------------------------------------------------------------------
TESSERACT_CANDIDATES = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
    os.path.expandvars(r"%LOCALAPPDATA%\Tesseract-OCR\tesseract.exe"),
]


def find_tesseract():
    exe = shutil.which("tesseract")
    if exe:
        return exe
    for c in TESSERACT_CANDIDATES:
        if c and os.path.exists(c):
            return c
    return None


def render_pdf_page(path, page_index, dpi=110):
    """Rend une page de PDF (ou une image) en image PIL — pour la visionneuse."""
    ext = os.path.splitext(path)[1].lower()
    if ext != ".pdf":
        with Image.open(path) as im:
            return ImageOps.exif_transpose(im).convert("RGB").copy()
    import pymupdf as fitz

    with fitz.open(path) as doc:
        page = doc[page_index]
        pix = page.get_pixmap(dpi=dpi)
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def page_count(path):
    if os.path.splitext(path)[1].lower() != ".pdf":
        return 1
    import pymupdf as fitz

    with fitz.open(path) as doc:
        return doc.page_count


def pdf_text(path):
    """Texte des pages (couche texte du PDF), liste par page."""
    if os.path.splitext(path)[1].lower() != ".pdf":
        return []
    import pymupdf as fitz

    with fitz.open(path) as doc:
        return [page.get_text() for page in doc]


# Texte imprimé des formulaires d'OM du bureau : sa présence ne signifie
# pas que les coordonnées (manuscrites) sont lisibles.
_FORM_MARKERS = ("Le(s) soussigné(s)", "Charge(nt) le", "ORDRE DE MISSION")


def ocr_image(pil_image, tesseract_exe, lang="eng", psm=6):
    with tempfile.TemporaryDirectory() as tmp:
        p = os.path.join(tmp, "page.png")
        pil_image.save(p)
        creation = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        res = subprocess.run(
            [tesseract_exe, p, "stdout", "-l", lang, "--psm", str(psm)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            creationflags=creation, timeout=180,
        )
        return res.stdout or ""


_MRZ_NAME_RE = re.compile(r"\b([A-Z]{2,}(?:<[A-Z]+)*)<<([A-Z]+)")


def parse_mrz_names(text):
    """Noms trouvés dans la zone codée des cartes d'identité (ligne
    « NOM<<PRENOM<AUTRES… »). Seuls les noms sont retenus — jamais les
    numéros de carte ou de registre national."""
    names = []
    for line in (text or "").splitlines():
        line = line.strip().replace(" ", "")
        if line.startswith(("IDBEL", "I<BEL", "P<")) or re.search(r"\d", line):
            # ligne 1 (n° de carte) ou 2 (dates) : ignorées
            if not line.startswith("P<"):
                continue
            line = line[5:]
        m = _MRZ_NAME_RE.search(line)
        if m:
            nom = m.group(1).replace("<", " ").strip()
            prenom = m.group(2).capitalize()
            entry = f"{nom} {prenom}"
            if entry not in names:
                names.append(entry)
    return names


def read_document(path, use_ocr=True, progress_cb=None, max_ocr_pages=5):
    """Lit un bail / OM. Retourne {"text_pages", "source", "mrz_names",
    "remarque"}. source = "texte" | "ocr" | "aucun"."""
    ext = os.path.splitext(path)[1].lower()
    out = {"text_pages": [], "source": "aucun", "mrz_names": [], "remarque": ""}
    if ext == ".docx":
        import docx

        d = docx.Document(path)
        out["text_pages"] = ["\n".join(p.text for p in d.paragraphs)]
        out["source"] = "texte"
    elif ext == ".pdf":
        pages = pdf_text(path)
        useful = [t for t in pages if len(t.strip()) > 40]
        is_form = useful and all(any(mk in t for mk in _FORM_MARKERS) for t in useful[:1])
        if useful and not is_form:
            out["text_pages"] = pages
            out["source"] = "texte"
        elif is_form:
            out["remarque"] = ("Formulaire du bureau rempli à la main : le texte imprimé est lisible, "
                               "mais pas les coordonnées manuscrites — à contrôler à l'écran.")
            out["text_pages"] = pages
    if (out["source"] == "aucun" and not out["remarque"]) or ext in (".jpg", ".jpeg", ".png"):
        exe = find_tesseract() if use_ocr else None
        if exe:
            n = min(page_count(path), max_ocr_pages)
            texts = []
            for i in range(n):
                if progress_cb:
                    progress_cb(f"Lecture (OCR) page {i + 1}/{n} de {os.path.basename(path)}…")
                img = render_pdf_page(path, i, dpi=300)
                texts.append(ocr_image(img, exe))
            out["text_pages"] = texts
            out["source"] = "ocr"
        elif not out["remarque"]:
            out["remarque"] = ("Document scanné : installez Tesseract (gratuit) pour lire le texte "
                               "et les cartes d'identité automatiquement.")
    out["mrz_names"] = parse_mrz_names("\n".join(out["text_pages"]))
    return out


def find_party_lines(text_pages):
    """Repère, dans un texte de bail, les lignes utiles autour de
    « bailleur », « preneur », « locataire », « domicilié »."""
    hits = []
    for pi, t in enumerate(text_pages):
        lines = [l.strip() for l in (t or "").splitlines()]
        for li, l in enumerate(lines):
            if re.search(r"\b(bailleur|preneur|locataire|domicili|demeurant|né le|nee le|née le)\b", l, re.I):
                ctx = " ".join(x for x in lines[li : li + 3] if x)
                if len(ctx) > 12 and ctx not in [h[1] for h in hits]:
                    hits.append((pi + 1, ctx[:220]))
    return hits[:40]


def compare_names(word_party_text, candidate_names):
    """Compare le nom d'une partie dans le Word aux noms lus (MRZ).
    Retourne (meilleur_nom, ratio)."""
    target = _norm(re.split(r",", word_party_text or "")[0])
    target = re.sub(r"^(M|MME|MR|MLLE|MELLE|MM|MMES)\.?\s+", "", target)
    best, best_r = None, 0.0
    for name in candidate_names:
        cand = _norm(name)
        r = difflib.SequenceMatcher(None, target, cand).ratio()
        # ordre prénom/nom inversé
        r2 = difflib.SequenceMatcher(None, " ".join(sorted(target.split())), " ".join(sorted(cand.split()))).ratio()
        r = max(r, r2)
        if r > best_r:
            best, best_r = name, r
    return best, best_r


# --------------------------------------------------------------------------
# Traitement par lots (plusieurs dossiers d'entrée d'un coup)
# --------------------------------------------------------------------------
def docx_needs_finalize(docx_path):
    """True si le Word n'a pas encore de photos sous le titre PHOTOGRAPHIES
    (lecture rapide du XML, sans ouvrir tout le document)."""
    import zipfile

    try:
        with zipfile.ZipFile(docx_path) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return False
    plain = re.sub(r"<[^>]+>", "", xml)
    if "PHOTOGRAPHIES" not in plain:
        return False
    i = xml.find("PHOTOGRAPHIES")
    j = xml.find("AUTRES PHOTOGRAPHIES", i + 1)
    segment = xml[i : j if j > 0 else len(xml)]
    return "<a:blip" not in segment


def find_pending(parent, initiales="GB", skipped=None):
    """Dossiers « ELLE » de l'expert `initiales` sous `parent` (le dossier
    lui-même, une tranche « 2026.3100 à 2026.3199 », ou « 2026 Dossiers
    DEF ») dont le Word reste à finaliser. Les dossiers des collègues (autres
    initiales) et ceux sans initiales sont ignorés et, si `skipped` est une
    liste, y sont ajoutés. Lecture seule. Retourne [{"folder", "docx",
    "nom_final"}]."""
    initiales = (initiales or "").strip().upper()
    candidates = []
    if securite.DOSSIER_RE.match(os.path.basename(os.path.normpath(parent))):
        candidates = [parent]
    else:
        for root, dirs, _files in os.walk(parent):
            depth = os.path.relpath(root, parent).count(os.sep)
            for d in sorted(dirs):
                if securite.DOSSIER_RE.match(d):
                    candidates.append(os.path.join(root, d))
            # ne descend pas dans les dossiers de dossiers, ni trop profond
            dirs[:] = [d for d in dirs if not securite.DOSSIER_RE.match(d)] if depth < 1 else []
    out = []
    for folder in candidates:
        info = parse_folder_name(folder)
        if info["type"].upper() not in ("ELLE", "ELE"):
            continue
        ini = folder_initials(folder)
        if initiales and ini != initiales:
            if skipped is not None:
                skipped.append((folder, ini or "sans initiales"))
            continue
        try:
            inv = scan_folder(folder)
        except Exception:  # noqa: BLE001
            continue
        for p in inv["docx"]:
            if docx_needs_finalize(p):
                out.append({"folder": folder, "docx": p, "nom_final": final_docx_name(folder)})
                break
    return out


def auto_parties(folder, docx_path, progress_cb=None):
    """Contrôle automatique (sans rien réécrire) : noms des parties du Word
    comparés aux noms lus sur les cartes d'identité du bail. Les noms
    « proches mais différents » sont surlignés en jaune. Retourne
    (updates, constats)."""
    inv = scan_folder(folder)
    parties = read_parties(docx_path)
    names, constats = [], []
    for doc in inv["documents"]:
        if "BAIL" not in _norm(os.path.basename(doc)):
            continue
        try:
            r = read_document(doc, progress_cb=progress_cb)
        except Exception as exc:  # noqa: BLE001
            constats.append(f"{os.path.basename(doc)} illisible ({exc})")
            continue
        names += [n for n in r["mrz_names"] if n not in names]
    updates = []
    if not names:
        constats.append("Pas de nom lu sur des cartes d'identité : parties à contrôler à la main.")
        return updates, constats
    for part in parties:
        name, ratio = compare_names(part["texte"], names)
        if ratio >= 0.99:
            constats.append(f"{part['role']} : identique à la carte d'identité ({name}).")
        elif ratio >= 0.6:
            updates.append({"index": part["index"], "texte": None, "surligner": True})
            constats.append(f"{part['role']} : « {part['texte'][:50]} » ≠ carte « {name} » → surligné en jaune.")
        else:
            constats.append(f"{part['role']} : aucun nom correspondant sur les cartes → à contrôler.")
    return updates, constats


def process_batch(jobs, date_cloture=None, tighten_layout=True, progress_cb=None, stop_flag=None):
    """Traite les dossiers l'un après l'autre. Une erreur sur un dossier
    n'arrête pas les suivants (et ne modifie pas son Word). `jobs` :
    [{"folder", "docx", "nom_final", "photos" (None = sélection
    automatique)}]. Retourne la liste des résultats."""
    results = []
    for k, job in enumerate(jobs, 1):
        if stop_flag and stop_flag():
            break
        name = os.path.basename(job["folder"])

        def say(msg, k=k, name=name):
            if progress_cb:
                progress_cb(k, len(jobs), name, msg)

        res = {"folder": job["folder"], "ok": False, "path": None, "constats": [], "erreur": None}
        try:
            say("Rangement des photos…")
            ranger_photos(job["folder"])
            photos = job.get("photos")
            if photos is None:
                items = classify_photos(scan_folder(job["folder"])["photo_images"])
                photos = [p for p, c, _n in items if c]
                exclues = [os.path.basename(p) for p, c, _n in items if not c]
                if exclues:
                    res["constats"].append("Photos exclues : " + ", ".join(exclues))
                res["constats"].append(
                    f"{len(photos)} photos injectées — photos d'index compteurs à retirer si présentes.")
            else:
                photos = [os.path.join(job["folder"], "Photos", os.path.basename(p))
                          if not os.path.exists(p) else p for p in photos]
            say("Contrôle des parties (bail)…")
            updates, constats = auto_parties(job["folder"], job["docx"], progress_cb=lambda m: say(m))
            res["constats"] += constats
            say("Finalisation du Word…")
            out = finalize(job["docx"], photos, date_cloture=date_cloture, parties_updates=updates,
                           rename_to=job.get("nom_final"), tighten_layout=tighten_layout,
                           progress_cb=lambda m: say(m))
            res.update(ok=True, path=out["path"], sauvegarde=out.get("sauvegarde"))
            res["constats"] += out["remarques"]
        except Exception as exc:  # noqa: BLE001
            res["erreur"] = str(exc)
            res["constats"] = [c for c in res["constats"] if "injectées" not in c]
        results.append(res)
    return results


def batch_report_text(results):
    lines = [f"Lot du {date.today().strftime('%d/%m/%Y')} — {sum(r['ok'] for r in results)}/{len(results)} dossier(s) finalisé(s)", ""]
    for r in results:
        lines.append(("✓ " if r["ok"] else "✗ ") + os.path.basename(r["folder"]))
        if r["erreur"]:
            lines.append("   ERREUR (Word non modifié) : " + r["erreur"])
        for c in r["constats"]:
            lines.append("   • " + c)
        lines.append("")
    return "\n".join(lines)


def keep_awake(on=True):
    """Empêche Windows de se mettre en veille pendant un lot."""
    try:
        import ctypes

        ES_CONTINUOUS, ES_SYSTEM_REQUIRED = 0x80000000, 0x00000001
        ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS | (ES_SYSTEM_REQUIRED if on else 0))
    except Exception:  # noqa: BLE001
        pass
