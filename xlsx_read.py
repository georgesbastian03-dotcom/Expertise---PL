"""Lecture du fichier Excel de mission via openpyxl."""
import openpyxl


def read_first_sheet(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows


def read_all_sheets(path: str):
    """Retourne la liste des lignes de TOUTES les feuilles du classeur (une
    liste de lignes par feuille), pour retrouver un petit tableau
    d'identification (bailleur, expert) où qu'il se trouve dans le fichier,
    sans dépendre de sa position exacte."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        all_rows = []
        for name in wb.sheetnames:
            ws = wb[name]
            all_rows.append([list(row) for row in ws.iter_rows(values_only=True)])
    finally:
        wb.close()
    return all_rows
